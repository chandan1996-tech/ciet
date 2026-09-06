from pptx import Presentation
from pptx.util import Inches, Pt

# Create a presentation object
prs = Presentation()

# Define a function to add a slide with title and bullet points
def add_bullet_slide(title, bullet_points):
    slide_layout = prs.slide_layouts[1]  # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title_placeholder = slide.shapes.title
    content_placeholder = slide.placeholders[1]

    title_placeholder.text = title
    content_placeholder.text = bullet_points[0]

    for point in bullet_points[1:]:
        p = content_placeholder.text_frame.add_paragraph()
        p.text = point
        p.level = 0

# Slides content
slides_content = [
    ("Solar System Components 101", ["Understanding the core parts of a solar PV system", "Source: GoGreenSolar"]),
    ("Introduction", ["Goal: Understand major components of a solar PV system", "Based on GoGreenSolar educational content"]),
    ("Overview of Core Components", [
        "Solar Panels (PV modules)",
        "Inverters",
        "Mounting & Racking System",
        "Monitoring System",
        "Optional: Batteries & Charge Controllers"
    ]),
    ("Solar Panels / PV Modules", [
        "Convert sunlight into DC electricity",
        "Types: Monocrystalline, Polycrystalline, Thin-film",
        "Efficiency typically 15–22%"
    ]),
    ("Inverters", [
        "Convert DC to AC electricity",
        "Types: String Inverters, Power Optimizers, Microinverters",
        "Trade-offs: shading resilience, cost, performance"
    ]),
    ("Mounting & Racking Systems", [
        "Secure panels to roof or ground",
        "Types: Fixed mounts, Tracking mounts",
        "Tilt & orientation affect efficiency"
    ]),
    ("Monitoring Systems", [
        "Track performance and system health",
        "Real-time dashboards and alerts",
        "Often integrated with inverters or as standalone tools"
    ]),
    ("Battery & Charge Control Systems", [
        "Store excess energy for later use",
        "Charge Controllers regulate charging (PWM or MPPT)",
        "Battery Management Systems (BMS) for monitoring"
    ]),
    ("Supporting Components", [
        "Net Meter: Measures energy flow to/from the grid",
        "Energy Management Systems (EMS): Optimize energy use"
    ]),
    ("Component Relationships & Flow", [
        "Flow: Panels → Inverter → Electrical Panel/Grid",
        "Optional: Battery & Monitoring overlay",
        "Configuration depends on system type"
    ]),
    ("GoGreenSolar Kit Options", [
        "Bundles include panels, inverters, racking, and more",
        "Options: Grid-tied, Off-grid, Battery-backed kits",
        "Warranties: 10–25 years depending on component"
    ]),
    ("Choosing the Right System", [
        "Consider space, shading, budget, and usage",
        "Panel type affects cost and performance",
        "Use tools to estimate and customize"
    ]),
    ("Summary & Q&A", [
        "Recap: Core components and functions",
        "Integrated design ensures efficiency",
        "Questions and further discussion"
    ])
]

# Add all slides
for title, content in slides_content:
    add_bullet_slide(title, content)

# Save presentation
prs.save("Solar_Components_101_Presentation.pptx")
